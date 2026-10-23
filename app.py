import os
import re
import json
from typing import Any
from urllib.parse import urljoin

import pandas as pd
import requests
import streamlit as st
from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document as DocxDocument
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

load_dotenv()

st.set_page_config(
    page_title="Enterprise RAG Document Assistant",
    page_icon="🔎",
    layout="wide",
)
st.title("🔎 Enterprise RAG Document Assistant")
st.caption(
    "Upload documents, retrieve relevant passages, and ask questions with source-grounded answers."
)

DEFAULT_OLLAMA_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")
DEFAULT_MODEL = os.getenv("OLLAMA_MODEL", "llama3.2")


def normalize_text(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def chunk_text(text: str, chunk_size: int = 900, overlap: int = 120) -> list[str]:
    """Split text into overlapping chunks while keeping the original reading order."""
    text = normalize_text(text)
    if not text:
        return []
    chunk_size = max(100, int(chunk_size))
    overlap = min(max(0, int(overlap)), chunk_size // 2)
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        if end < len(text):
            # Prefer a word boundary when it is reasonably close to the target.
            boundary = text.rfind(" ", start + int(chunk_size * 0.75), end)
            if boundary > start:
                end = boundary
        part = text[start:end].strip()
        if part:
            chunks.append(part)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def read_uploaded_file(uploaded_file) -> tuple[str, str]:
    """Return extracted text and a short description for the supported upload type."""
    name = uploaded_file.name
    suffix = Path(name).suffix.lower()
    raw = uploaded_file.getvalue()

    if suffix == ".txt":
        return raw.decode("utf-8", errors="replace"), "TXT"
    if suffix == ".pdf":
        import io
        reader = PdfReader(io.BytesIO(raw))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages.append(f"[Page {page_number}] {page_text}")
        return "\n".join(pages), "PDF"
    if suffix == ".docx":
        import io
        doc = DocxDocument(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts), "DOCX"
    if suffix == ".csv":
        import io
        df = pd.read_csv(io.BytesIO(raw)).fillna("")
        # Turn each row into a readable record to keep table context during retrieval.
        rows = []
        for idx, row in df.iterrows():
            fields = [f"{col}: {row[col]}" for col in df.columns if str(row[col]).strip()]
            if fields:
                rows.append(f"CSV row {idx + 2}: " + "; ".join(fields))
        return "\n".join(rows), "CSV"
    raise ValueError(f"Unsupported file type: {suffix}")


def build_chunks(uploaded_files, chunk_size: int, overlap: int) -> list[dict[str, Any]]:
    documents = []
    for uploaded in uploaded_files:
        text, file_type = read_uploaded_file(uploaded)
        parts = chunk_text(text, chunk_size=chunk_size, overlap=overlap)
        for chunk_index, part in enumerate(parts, start=1):
            documents.append({
                "source": uploaded.name,
                "file_type": file_type,
                "chunk_id": chunk_index,
                "text": part,
            })
    return documents


def retrieve(question: str, chunks: list[dict[str, Any]], top_k: int = 4):
    """Local TF-IDF retrieval; simple and private, with no external embedding API."""
    if not chunks or not question.strip():
        return []
    corpus = [item["text"] for item in chunks]
    vectorizer = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), max_features=50000)
    try:
        matrix = vectorizer.fit_transform(corpus + [question])
    except ValueError:
        return []
    similarities = cosine_similarity(matrix[-1], matrix[:-1]).flatten()
    ranked = similarities.argsort()[::-1][:max(1, min(int(top_k), len(chunks)))]
    results = []
    for idx in ranked:
        score = float(similarities[idx])
        if score > 0:
            item = dict(chunks[int(idx)])
            item["score"] = score
            results.append(item)
    return results


def call_ollama(question: str, context_items: list[dict[str, Any]], model: str, base_url: str) -> str:
    context = "\n\n".join(
        f"[Source: {item['source']} | chunk {item['chunk_id']}]\n{item['text']}"
        for item in context_items
    )
    prompt = f"""You are a careful document assistant. Answer the user's question using only the context below.
If the context does not contain the answer, say that the uploaded documents do not provide enough information.
Do not invent facts. Cite supporting passages inline using [filename, chunk N].
Keep the answer concise and distinguish facts from reasonable interpretation.

CONTEXT:
{context}

QUESTION:
{question}

ANSWER:"""
    url = base_url.rstrip("/") + "/api/generate"
    response = requests.post(
        url,
        json={"model": model, "prompt": prompt, "stream": False},
        timeout=120,
    )
    response.raise_for_status()
    payload = response.json()
    return str(payload.get("response", "")).strip() or "The model returned an empty response."


def ollama_available(base_url: str) -> bool:
    try:
        response = requests.get(base_url.rstrip("/") + "/api/tags", timeout=3)
        return response.ok
    except requests.RequestException:
        return False


with st.sidebar:
    st.header("Settings")
    chunk_size = st.slider("Chunk size (characters)", min_value=300, max_value=1800, value=900, step=100)
    overlap = st.slider("Chunk overlap (characters)", min_value=0, max_value=300, value=120, step=20)
    top_k = st.slider("Retrieved passages", min_value=1, max_value=8, value=4)
    st.divider()
    st.subheader("Local LLM (Ollama)")
    ollama_url = st.text_input("Ollama base URL", value=DEFAULT_OLLAMA_URL)
    ollama_model = st.text_input("Model name", value=DEFAULT_MODEL)
    use_llm = st.checkbox("Generate answer with Ollama", value=True)
    st.caption("Ollama runs locally. If unavailable, you can still inspect retrieved passages.")

uploaded_files = st.file_uploader(
    "Upload documents",
    type=["pdf", "docx", "txt", "csv"],
    accept_multiple_files=True,
    help="Supported: PDF, DOCX, TXT, CSV. Scanned PDFs may need OCR before upload.",
)

if "chunks" not in st.session_state:
    st.session_state.chunks = []
if "indexed_names" not in st.session_state:
    st.session_state.indexed_names = []
if "chat_history" not in st.session_state:
    st.session_state.chat_history = []

col1, col2 = st.columns([1, 1])
with col1:
    build_clicked = st.button("Build / Refresh document index", type="primary", use_container_width=True)
with col2:
    if st.button("Clear chat history", use_container_width=True):
        st.session_state.chat_history = []
        st.rerun()

if build_clicked:
    if not uploaded_files:
        st.warning("Upload at least one document first.")
    else:
        try:
            with st.spinner("Reading documents and creating searchable chunks..."):
                st.session_state.chunks = build_chunks(uploaded_files, chunk_size, overlap)
                st.session_state.indexed_names = [f.name for f in uploaded_files]
            st.session_state.chat_history = []
            st.success(
                f"Indexed {len(st.session_state.chunks)} passages from "
                f"{len(st.session_state.indexed_names)} file(s)."
            )
        except Exception as exc:
            st.error(f"Could not process uploads: {type(exc).__name__}: {exc}")

if st.session_state.chunks:
    st.caption("Indexed files: " + ", ".join(st.session_state.indexed_names))
    question = st.chat_input("Ask a question about your documents")
    for message in st.session_state.chat_history:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if question:
        st.session_state.chat_history.append({"role": "user", "content": question})
        with st.chat_message("user"):
            st.markdown(question)

        retrieved = retrieve(question, st.session_state.chunks, top_k=top_k)
        with st.chat_message("assistant"):
            if not retrieved:
                answer = "I couldn't find a relevant passage in the indexed documents. Try rephrasing the question."
                st.markdown(answer)
            elif use_llm and ollama_available(ollama_url):
                try:
                    with st.spinner(f"Generating a grounded answer with {ollama_model}..."):
                        answer = call_ollama(question, retrieved, ollama_model, ollama_url)
                    st.markdown(answer)
                except requests.RequestException as exc:
                    answer = (
                        "I found relevant passages, but Ollama could not generate an answer. "
                        "Check that Ollama is running and the model is available. "
                        f"Technical detail: {type(exc).__name__}"
                    )
                    st.warning(answer)
            elif use_llm:
                answer = (
                    "Ollama is not reachable. Start Ollama and make sure the selected model is installed. "
                    "Relevant passages are shown below."
                )
                st.info(answer)
            else:
                answer = "Retrieved relevant passages below. Enable Ollama to generate a grounded answer."
                st.markdown(answer)

            with st.expander("Retrieved sources", expanded=True):
                for item in retrieved:
                    st.markdown(
                        f"**{item['source']} · chunk {item['chunk_id']}** "
                        f"(similarity: {item['score']:.3f})"
                    )
                    st.write(item["text"])
                    st.divider()

        st.session_state.chat_history.append({"role": "assistant", "content": answer})
else:
    st.info("Upload one or more files, then select **Build / Refresh document index** to start.")

with st.expander("How this starter works"):
    st.markdown(
        """
        1. **Ingestion:** extracts text from PDF, DOCX, TXT, and CSV.
        2. **Chunking:** splits extracted text into overlapping passages.
        3. **Retrieval:** uses local TF-IDF similarity to find relevant passages.
        4. **Generation:** optionally sends the retrieved context to a local Ollama model.
        5. **Grounding:** the prompt asks the model to stay within the supplied context and cite source chunks.

        This starter uses TF-IDF rather than an embedding-based vector database. For a stronger portfolio version,
        add embeddings + Chroma/FAISS, automated retrieval evaluation, OCR for scanned PDFs, authentication,
        access controls, and tests before describing it as production-ready.
        """
    )
